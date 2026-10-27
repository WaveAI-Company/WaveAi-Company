"""Preenche o template AMS Conecta com o conteudo do WaveAI.

O template e a identidade visual: moldura do Governo de SP nos slides de
conteudo, capa e encerramento escuros com o logo AMS Conecta. Aqui so entra
texto e elementos no espaco vazio da direita -- nada da moldura e tocado.

Precisa de `template.pptx` (o modelo da organizacao) na mesma pasta. Rodar:
`python gerar_deck.py` -> WaveAI_Pitch_Conecta.pptx
"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# --- marca WaveAI (Design/logos_icones) -------------------------------------
NAVY = RGBColor(0x0B, 0x12, 0x20)
AZUL = RGBColor(0x2A, 0x5B, 0xC7)
TEAL = RGBColor(0x4F, 0xD1, 0xC5)
TEAL_ESC = RGBColor(0x0F, 0x7A, 0x70)
CINZA = RGBColor(0x30, 0x30, 0x30)       # cor do corpo no template
CINZA_CL = RGBColor(0x6B, 0x72, 0x80)
GELO = RGBColor(0xF5, 0xF7, 0xFA)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)

F_TIT = "Verdana"              # fonte dos titulos no template
F_TXT = "Google Sans Text"     # fonte do corpo no template

# Area util, em polegadas num slide de 13,33 x 7,5. A moldura do template tem
# cabecalho do governo no topo e faixas pretas a partir de ~6,55 no pe: nada de
# conteudo abaixo disso. O template so ocupa a esquerda; a direita fica vazia.
COL_E, LARG_E = 0.96, 6.70     # coluna esquerda (texto), como no template
COL_D, LARG_D = 8.00, 4.40     # coluna direita (vazia no template)


def limpa(shape):
    """Tira contorno e sombra herdados -- o padrao do PowerPoint e feio."""
    shape.line.fill.background()
    shape.shadow.inherit = False


def texto(slide, x, y, w, h, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    cx = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = cx.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tf


def run(par, txt, tam, cor, negrito=False, fonte=F_TXT, espaco_antes=None):
    r = par.add_run()
    r.text = txt
    r.font.size = Pt(tam)
    r.font.bold = negrito
    r.font.name = fonte
    r.font.color.rgb = cor
    if espaco_antes is not None:
        par.space_before = Pt(espaco_antes)
    return r


def paragrafo(tf, primeiro=False):
    return tf.paragraphs[0] if primeiro else tf.add_paragraph()


def corpo(tf, linhas, tam=14, cor=CINZA, espaco=8):
    """linhas: lista de (texto, negrito) ou texto puro."""
    for i, ln in enumerate(linhas):
        txt, neg = ln if isinstance(ln, tuple) else (ln, False)
        p = paragrafo(tf, i == 0)
        p.line_spacing = 1.22
        run(p, txt, tam, cor, neg, espaco_antes=(0 if i == 0 else espaco))


def troca_titulo(slide, novo):
    """Reescreve o titulo mantendo Verdana 44pt do template.

    A caixa do template tem 6,7-7,8 pol e ancora vertical centralizada: titulo
    de duas palavras longas QUEBRA em duas linhas e, por estar centralizado,
    desce em cima do corpo. Medido no render: "Problema identificado",
    "Proposta de solucao" e "Diferenciais competitivos" colidiam. Alargar a
    caixa ate a margem direita resolve sem mexer em fonte nem tamanho.
    """
    for sh in slide.shapes:
        if sh.has_text_frame and sh.name.startswith(("Título", "Titulo")):
            p = sh.text_frame.paragraphs[0]
            p.runs[0].text = novo
            for r in p.runs[1:]:
                r.text = ""
            sh.left, sh.width = Inches(COL_E - 0.01), Inches(11.43)
            return sh
    raise SystemExit("titulo nao encontrado")


def caixa_corpo(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.name.startswith("CaixaDeTexto"):
            return sh
    raise SystemExit("caixa de corpo nao encontrada")


def limpa_corpo(slide):
    """Esvazia a caixa de corpo do template e devolve o text_frame."""
    sh = caixa_corpo(slide)
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    return sh, tf


def cartao(slide, x, y, w, h, fundo, borda=None):
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    c.adjustments[0] = 0.08
    c.fill.solid()
    c.fill.fore_color.rgb = fundo
    c.shadow.inherit = False
    if borda:
        c.line.color.rgb = borda
        c.line.width = Pt(1.25)
    else:
        c.line.fill.background()
    c.text_frame.text = ""
    return c


def estatistica(slide, x, y, w, numero, rotulo, cor_num=AZUL, fonte_h=1.30):
    """Numero grande + rotulo pequeno, sem moldura."""
    tf = texto(slide, x, y, w, fonte_h * 0.62)
    run(tf.paragraphs[0], numero, 34, cor_num, True, F_TIT)
    tl = texto(slide, x, y + fonte_h * 0.56, w, fonte_h * 0.46)
    p = tl.paragraphs[0]
    p.line_spacing = 1.15
    run(p, rotulo, 10.5, CINZA_CL, False)


pres = Presentation("template.pptx")
s = pres.slides

# ============================================================ SLIDE 1 — capa
cap = s[0]
for sh in cap.shapes:
    if sh.has_text_frame and "Nome do projeto" in sh.text_frame.text:
        p = sh.text_frame.paragraphs[0]
        p.runs[0].text = "WaveAI"
        for r in p.runs[1:]:
            r.text = ""
        p.runs[0].font.size = Pt(60)
        p.runs[0].font.bold = True
        p.runs[0].font.name = F_TIT
        p.runs[0].font.color.rgb = BRANCO
        p.alignment = PP_ALIGN.CENTER
        break
tf = texto(cap, 0.51, 4.30, 12.31, 0.55, PP_ALIGN.CENTER)
run(tf.paragraphs[0], "Sua mente em ondas. Seu bem-estar em movimento.",
    19, TEAL, False, F_TXT)

# =============================================== SLIDE 2 — identificação
ident = s[1]
for sh in ident.shapes:
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text
    if "Nome do Projeto" in t:
        p = sh.text_frame.paragraphs[0]
        p.runs[0].text = "WaveAI"
        p.runs[0].font.size = Pt(54)
        p.runs[1].text = "\nSua mente em ondas. Seu bem-estar em movimento."
        p.runs[1].font.size = Pt(21)
        p.runs[1].font.bold = False
        p.runs[1].font.color.rgb = TEAL_ESC
    elif "UNIDADE ESCOLAR" in t:
        ps = sh.text_frame.paragraphs
        ps[0].runs[0].text = "FATEC ZONA LESTE"
        ps[0].runs[0].font.bold = True
        ps[0].runs[0].font.size = Pt(16)
        ps[1].runs[0].text = ("Tiago Bryan Ramos de Oliveira  ·  "
                              "Gustavo Mendes Vetieri Mariano  ·  "
                              "Nickolas Maia de Araujo")
        ps[1].runs[0].font.size = Pt(13)
        ps[1].runs[0].font.color.rgb = CINZA

# ==================================================== SLIDE 3 — o problema
p3 = s[2]
troca_titulo(p3, "Problema identificado")
sh, tf = limpa_corpo(p3)
sh.top, sh.height = Inches(2.30), Inches(3.9)
sh.width = Inches(LARG_E)
corpo(tf, [
    ("O mercado: bem-estar mental deixou de ser pauta de RH e virou "
     "obrigação legal.", True),
    "Desde 26 de maio de 2026, a NR-1 obriga toda empresa brasileira a "
    "identificar e gerenciar riscos psicossociais no PGR — com fiscalização "
    "e multa. Apenas 5% dos RHs se dizem preparados.",
    ("A dor: ninguém tem com o que medir.", True),
    "A gestão de bem-estar se apoia em questionário e autorrelato — escalas "
    "de 1 a 5 respondidas justamente por quem está sobrecarregado. Não existe "
    "sinal objetivo na conta.",
    ("O público: empresas obrigadas pela NR-1, profissionais de bem-estar que "
     "acompanham pessoas e adultos que já medem passos, sono e batimentos — "
     "e não medem o cérebro.", True),
])
estatistica(p3, COL_D, 2.30, LARG_D, "393 mil",
            "afastamentos do INSS por transtornos mentais em 2025 — "
            "79% acima de 2023")
estatistica(p3, COL_D, 3.72, LARG_D, "3×",
            "foi o crescimento dos afastamentos por burnout entre 2023 e 2025",
            TEAL_ESC)
estatistica(p3, COL_D, 5.14, LARG_D, "R$ 954 mi",
            "o custo desses benefícios para o INSS no último ano", NAVY)
tfn = texto(p3, COL_D, 6.18, LARG_D, 0.30)
run(tfn.paragraphs[0], "Fonte: ANAMT, com microdados oficiais do INSS (2026)",
    8.5, CINZA_CL)

# ================================================== SLIDE 4 — a solução
p4 = s[3]
troca_titulo(p4, "Proposta de solução")
sh, tf = limpa_corpo(p4)
sh.top, sh.height = Inches(2.30), Inches(3.9)
sh.width = Inches(LARG_E)
corpo(tf, [
    ("O WaveAI tira o eletroencefalograma do hospital e o põe na rotina.", True),
    "Um aparelho de EEG de consumo lê o sinal elétrico do córtex frontal a "
    "512 amostras por segundo. O aplicativo capta, o servidor analisa e a "
    "pessoa passa a ver a tendência do próprio sinal ao longo das semanas.",
    "A transformação prática: onde antes havia uma resposta de questionário, "
    "passa a haver uma série temporal medida — com indicador de qualidade do "
    "sinal em cada sessão.",
    ("Ferramenta não-clínica e não-diagnóstica.", True),
    "O WaveAI mede, registra e organiza. Quem interpreta é o profissional de "
    "bem-estar, que agora decide com número e não apenas com relato.",
])
passos = [
    ("1", "Captar", "Aparelho de consumo, 512 Hz, sinal bruto"),
    ("2", "Conferir", "Protocolo guiado verifica o instrumento antes de medir"),
    ("3", "Analisar", "13 indicadores, com qualidade de sinal por sessão"),
    ("4", "Acompanhar", "A tendência ao longo do tempo é o produto"),
]
y = 2.30
for num, titulo, desc in passos:
    cartao(p4, COL_D, y, LARG_D, 0.92, GELO)
    tb = texto(p4, COL_D + 0.20, y + 0.16, 0.42, 0.60)
    run(tb.paragraphs[0], num, 22, TEAL_ESC, True, F_TIT)
    tt = texto(p4, COL_D + 0.68, y + 0.14, LARG_D - 0.88, 0.62)
    run(tt.paragraphs[0], titulo, 13, NAVY, True)
    pd = tt.add_paragraph()
    pd.line_spacing = 1.1
    run(pd, desc, 10, CINZA_CL, espaco_antes=2)
    y += 1.03

# =============================================== SLIDE 5 — diferenciais
p5 = s[4]
troca_titulo(p5, "Diferenciais competitivos")
sh, tf = limpa_corpo(p5)
sh.left, sh.top, sh.height = Inches(COL_E), Inches(2.30), Inches(3.9)
sh.width = Inches(LARG_E)
corpo(tf, [
    ("Contra o questionário, que é a alternativa real: medimos.", True),
    "As plataformas de bem-estar corporativo entregam formulário, conteúdo e "
    "sessão de terapia. Nenhuma entrega sinal fisiológico. É outra categoria "
    "de evidência, e é a que a NR-1 vai cobrar.",
    ("Contra o EEG clínico: cabe na rotina.", True),
    "Equipamento de laboratório exige sala, técnico e agendamento, e produz "
    "uma fotografia isolada. Nós produzimos série contínua, em casa, sem "
    "operador — que é o que revela tendência.",
    ("Contra quem fizer igual: já estamos no ar.", True),
    "A plataforma está em produção e acessível pelo navegador, construída por "
    "três estudantes, com infraestrutura que custa praticamente zero em repouso.",
])
selos = [
    ("A tela nunca mente",
     "Se a medida saiu ruim, está escrito que saiu ruim. “Não foi medido "
     "direito” e “foi medido e não deu” são frases diferentes."),
    ("O dado é do titular",
     "Cifrado, com leitura auditada e exclusão imediata. O compartilhamento "
     "com o profissional nasce desligado — quem liga é a pessoa."),
    ("Análise rastreável",
     "Cada resultado grava a versão do motor que o produziu. Dá para auditar "
     "como aquele número foi obtido, meses depois."),
]
y = 2.30
for titulo, desc in selos:
    cartao(p5, COL_D, y, LARG_D, 1.22, BRANCO, TEAL)
    tt = texto(p5, COL_D + 0.22, y + 0.17, LARG_D - 0.44, 0.90)
    run(tt.paragraphs[0], titulo, 12.5, TEAL_ESC, True)
    pd = tt.add_paragraph()
    pd.line_spacing = 1.14
    run(pd, desc, 9.5, CINZA, espaco_antes=3)
    y += 1.33

# ============================================ SLIDE 6 — modelo de negócios
p6 = s[5]
troca_titulo(p6, "Modelo de Negócios")
sh, tf = limpa_corpo(p6)
sh.left, sh.top, sh.height = Inches(COL_E), Inches(2.30), Inches(3.9)
sh.width = Inches(LARG_E)
corpo(tf, [
    ("Três receitas, uma porta de entrada.", True),
    "1 · Assinatura do profissional de bem-estar, por paciente acompanhado. "
    "É o canal: ele traz o paciente junto.",
    "2 · Pacote corporativo para a NR-1, cobrado por colaborador aderente. "
    "A empresa recebe indicador agregado e anônimo, com adesão voluntária — "
    "o empregador nunca vê o dado de um indivíduo. É a linha que escala.",
    "3 · Licença do motor de análise como API, para quem já tem aparelho e "
    "não tem a inteligência.",
])
# A projecao de 24 meses saiu a pedido do fundador: as premissas de preco nao
# foram validadas com ninguem do setor, e numero inventado e a parte mais facil
# de um pitch ser derrubada. O dado de mercado fica -- esse tem fonte.
cartao(p6, COL_D, 2.30, LARG_D, 2.05, NAVY)
tm = texto(p6, COL_D + 0.30, 2.60, LARG_D - 0.60, 1.50)
run(tm.paragraphs[0], "US$ 2,7 bi", 38, TEAL, True, F_TIT)
pm = tm.add_paragraph()
pm.line_spacing = 1.22
run(pm, "é o tamanho projetado do mercado global de EEG vestível em 2030, "
        "contra US$ 1,5 bi em 2025 — crescimento de 12,4% ao ano",
    11, GELO, espaco_antes=8)

tp = texto(p6, COL_D, 4.70, LARG_D, 0.28)
run(tp.paragraphs[0], "PARA SAIR DO PILOTO", 10, CINZA_CL, True)

pedidos = [
    ("Laboratório com EEG de referência",
     "fecha o estudo de fidelidade de sinal"),
    ("Parecer regulatório e revisão jurídica",
     "define até onde podemos afirmar"),
    ("Piloto pago com profissionais reais",
     "primeiro dado de uso e validação de preço"),
]
y = 5.06
for titulo, porque in pedidos:
    tl = texto(p6, COL_D + 0.22, y, LARG_D - 0.22, 0.48)
    run(tl.paragraphs[0], titulo, 11, NAVY, True)
    pd = tl.add_paragraph()
    pd.line_spacing = 1.1
    run(pd, porque, 9, CINZA_CL, espaco_antes=1)
    mk = p6.shapes.add_shape(MSO_SHAPE.OVAL, Inches(COL_D + 0.02),
                             Inches(y + 0.07), Inches(0.09), Inches(0.09))
    mk.fill.solid()
    mk.fill.fore_color.rgb = TEAL
    limpa(mk)
    y += 0.50

# ================================================== SLIDE 7 — encerramento
# O brasao do Governo de SP vem na imagem de fundo e termina por volta de
# 4,2 pol -- medido no render. Tudo abaixo disso, ou colide com ele.
p7 = s[6]
for sh in p7.shapes:
    if sh.has_text_frame and "OBRIGADO" in sh.text_frame.text:
        sh.left, sh.width = Inches(0.95), Inches(11.43)
        sh.top, sh.height = Inches(4.42), Inches(0.62)
        p = sh.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].text = "Sua mente em ondas."
        p.runs[0].font.size = Pt(32)
        p.runs[0].font.name = F_TIT
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = NAVY
        break
tf = texto(p7, 0.95, 5.06, 11.43, 0.62, PP_ALIGN.CENTER)
run(tf.paragraphs[0], "Seu bem-estar em movimento.", 32, TEAL_ESC, True, F_TIT)
tf2 = texto(p7, 0.95, 5.92, 11.43, 0.30, PP_ALIGN.CENTER)
run(tf2.paragraphs[0], "waveai.tec.br", 15, AZUL, True)
tf3 = texto(p7, 0.95, 6.28, 11.43, 0.28, PP_ALIGN.CENTER)
run(tf3.paragraphs[0],
    "Tiago Bryan Ramos de Oliveira  ·  Gustavo Mendes Vetieri Mariano  ·  "
    "Nickolas Maia de Araujo", 10, CINZA_CL)

pres.save("WaveAI_Pitch_Conecta.pptx")
print("gerado: WaveAI_Pitch_Conecta.pptx")
